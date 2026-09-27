"""
Unione e caricamento della configurazione completa.

Contiene:
- `_festa_da_dict`: costruisce un oggetto Festa da un dict YAML
- `_ricorrenza_da_dict`: costruisce un oggetto Ricorrenza da un dict YAML
- `unisci_config`: unisce dati YAML + Excel in un oggetto Config
- `carica_config`: API pubblica che legge tutto e restituisce Config

Uso tipico:
    from config import carica_config

    config = carica_config("san_pietro_in_silki")
    print(config.anno)             # 2027
    print(config.periodi[0].dal)   # date(2027, 10, 1)
"""

from __future__ import annotations

import logging
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from dataclass_config import (
    Config,
    Festa,
    Intenzione,
    Matrimonio,
    Periodo,
    Ricorrenza,
)
from lettura_excel import (
    leggi_foglio_impostazioni,
    leggi_foglio_intenzioni,
    leggi_foglio_matrimoni,
    leggi_foglio_note,
)
from lettura_yaml import leggi_yaml

logger = logging.getLogger(__name__)

# ======================================================================
# UNIONE CONFIGURAZIONE
# ======================================================================


def _festa_da_dict(dati: dict) -> Festa:
    """Costruisce un oggetto Festa da un dict del YAML."""
    return Festa(
        nome=dati.get("nome", ""),
        mese=dati.get("mese"),
        giorno=dati.get("giorno"),
        tipo=dati.get("tipo", "fissa"),
        salta_se_domenica=dati.get("salta_se_domenica", False),
        orari_ridotti=dati.get("orari_ridotti"),
    )


def _ricorrenza_da_dict(dati: dict, anno: int) -> Ricorrenza | None:
    """Costruisce un oggetto Ricorrenza da un dict del YAML.

    Ritorna None se la data non può essere determinata (formato non
    supportato in questa fase).
    """
    nome = dati.get("nome", "")
    data_spec = dati.get("data", {})

    # Caso 1: data fissa (mese + giorno, opzionale durata_giorni)
    if "mese" in data_spec and "giorno" in data_spec:
        data_inizio = date(anno, data_spec["mese"], data_spec["giorno"])
        durata = dati.get("durata_giorni", 1)
        data_fine = data_inizio + timedelta(days=durata - 1) if durata > 1 else None
        return Ricorrenza(
            nome=nome,
            data_inizio=data_inizio,
            data_fine=data_fine,
            tipo=dati.get("tipo", "semplice"),
            compilazione=dati.get("compilazione", "auto"),
            nota=dati.get("nota", ""),
        )

    # Caso 2: intervallo (giorno_da + giorno_a, stesso mese o mese_a)
    if "giorno_da" in data_spec and "giorno_a" in data_spec:
        mese = data_spec.get("mese")
        mese_a = data_spec.get("mese_a", mese)
        data_inizio = date(anno, mese, data_spec["giorno_da"])
        data_fine = date(anno, mese_a, data_spec["giorno_a"])
        return Ricorrenza(
            nome=nome,
            data_inizio=data_inizio,
            data_fine=data_fine,
            tipo=dati.get("tipo", "semplice"),
            compilazione=dati.get("compilazione", "auto"),
            nota=dati.get("nota", ""),
        )

    # Caso 3: tipo speciale (es. ultima_domenica_maggio) → per ora None
    if "tipo" in data_spec:
        logger.info(
            "Ricorrenza '%s' con tipo speciale '%s' non ancora supportata. Ignorata.",
            nome,
            data_spec["tipo"],
        )
        return None

    logger.warning(
        "Ricorrenza '%s' ha un formato data non riconosciuto. Ignorata.",
        nome,
    )
    return None


def unisci_config(dati_yaml: dict, dati_excel: dict) -> Config:
    """Unisce i dati da YAML (stabili) e da Excel (annuali) in un oggetto Config.

    Args:
        dati_yaml: dizionario dalle regole.yaml
        dati_excel: dizionario con anno, nome_parrocchia, citta, periodi,
                    intenzioni, matrimoni, note

    Returns:
        Oggetto Config completo.
    """
    anno = dati_excel.get("anno", 0)

    # ------------------------------------------------------------------
    # Nome parrocchia e città: YAML vince (dato ufficiale)
    # ------------------------------------------------------------------
    nome_yaml = dati_yaml.get("nome_parrocchia", "")
    nome_excel = dati_excel.get("nome_parrocchia", "")

    if nome_yaml and nome_excel and nome_yaml != nome_excel:
        logger.warning(
            "Mismatch 'nome_parrocchia': YAML='%s', Excel='%s'. Uso YAML.",
            nome_yaml,
            nome_excel,
        )
    nome_parrocchia = nome_yaml or nome_excel

    citta_yaml = dati_yaml.get("citta", "")
    citta_excel = dati_excel.get("citta", "")
    if citta_yaml and citta_excel and citta_yaml != citta_excel:
        logger.warning(
            "Mismatch 'citta': YAML='%s', Excel='%s'. Uso YAML.",
            citta_yaml,
            citta_excel,
        )
    citta = citta_yaml or citta_excel

    # ------------------------------------------------------------------
    # Anno: Excel vince (dato annuale)
    # ------------------------------------------------------------------
    # (per ora non c'è un campo anno nel YAML, ma teniamo il check per il futuro)
    anno_yaml = dati_yaml.get("anno")
    if anno_yaml and anno_yaml != anno:
        logger.warning(
            "Mismatch 'anno': YAML=%s, Excel=%s. Uso Excel.",
            anno_yaml,
            anno,
        )

    # ------------------------------------------------------------------
    # Feste (da YAML)
    # ------------------------------------------------------------------
    festivi_fissi = [_festa_da_dict(f) for f in dati_yaml.get("festivi_fissi", [])]
    eccezioni_festivi = [_festa_da_dict(f) for f in dati_yaml.get("eccezioni_festivi", [])]

    # ------------------------------------------------------------------
    # Divieti pomeridiani (da YAML)
    # ------------------------------------------------------------------
    divieti = dati_yaml.get("divieti_pomeridiani", {})
    divieti_fissi = [_festa_da_dict(f) for f in divieti.get("fissi", [])]
    divieti_mobili = divieti.get("mobili", [])

    # ------------------------------------------------------------------
    # Ricorrenze proprie (da YAML)
    # ------------------------------------------------------------------
    ricorrenze: list[Ricorrenza] = []
    for r_dict in dati_yaml.get("ricorrenze_proprie", []):
        ric = _ricorrenza_da_dict(r_dict, anno)
        if ric is not None:
            ricorrenze.append(ric)

    # ------------------------------------------------------------------
    # Periodi (da Excel)
    # ------------------------------------------------------------------
    periodi = [
        Periodo(
            dal=p["dal"],
            al=p["al"],
            orari_feriali=p["orari_feriali"],
            orari_festivi=p["orari_festivi"],
        )
        for p in dati_excel.get("periodi", [])
    ]

    # ------------------------------------------------------------------
    # Costruisci Config
    # ------------------------------------------------------------------
    return Config(
        # Da YAML
        nome_parrocchia=nome_parrocchia,
        citta=citta,
        festivi_fissi=festivi_fissi,
        eccezioni_festivi=eccezioni_festivi,
        divieti_pomeridiani_fissi=divieti_fissi,
        divieti_pomeridiani_mobili=divieti_mobili,
        ricorrenze_proprie=ricorrenze,
        celebrazioni_mobili=dati_yaml.get("celebrazioni_mobili", []),
        giorni_accorpati=dati_yaml.get("giorni_accorpati", []),
        nove_mercoledi=dati_yaml.get("nove_mercoledi", {}),
        matrimoni_config=dati_yaml.get("matrimoni", {}),
        intenzioni_config=dati_yaml.get("intenzioni", {}),
        # Da Excel
        anno=anno,
        periodi=periodi,
        intenzioni=dati_excel.get("intenzioni", []),
        matrimoni_prenotati=dati_excel.get("matrimoni", []),
        note_annuali=dati_excel.get("note", []),
    )


# ======================================================================
# API PUBBLICA
# ======================================================================


def carica_config(
    profilo: str,
    cartella_configs: Path | None = None,
) -> Config:
    """Carica la configurazione completa di un profilo.

    Legge:
    - configs/<profilo>/regole.yaml
    - configs/<profilo>/config.xlsx

    Unisce i dati e restituisce un oggetto Config.

    Args:
        profilo: nome del profilo (es. "san_pietro_in_silki")
        cartella_configs: percorso alla cartella "configs" (opzionale).
            Se None, usa il default: <root_progetto>/configs

    Returns:
        Oggetto Config completo.

    Raises:
        FileNotFoundError: se il profilo o i file non esistono
        ValueError: se i file sono malformati
    """
    # ------------------------------------------------------------------
    # Percorsi
    # ------------------------------------------------------------------
    if cartella_configs is None:
        root_progetto = Path(__file__).resolve().parent.parent
        cartella_configs = root_progetto / "configs"

    profilo_dir = cartella_configs / profilo

    if not profilo_dir.exists():
        raise FileNotFoundError(f"Profilo non trovato: {profilo_dir}")

    percorso_yaml = profilo_dir / "regole.yaml"
    percorso_excel = profilo_dir / "config.xlsx"

    if not percorso_yaml.exists():
        raise FileNotFoundError(f"File regole.yaml non trovato: {percorso_yaml}")

    if not percorso_excel.exists():
        raise FileNotFoundError(f"File config.xlsx non trovato: {percorso_excel}")

    # ------------------------------------------------------------------
    # Lettura YAML
    # ------------------------------------------------------------------
    logger.info("Lettura YAML: %s", percorso_yaml)
    dati_yaml = leggi_yaml(percorso_yaml)

    # ------------------------------------------------------------------
    # Lettura Excel
    # ------------------------------------------------------------------
    logger.info("Lettura Excel: %s", percorso_excel)
    wb = load_workbook(percorso_excel, data_only=True)

    # Foglio Impostazioni
    if "Impostazioni" not in wb.sheetnames:
        raise ValueError(f"Foglio 'Impostazioni' mancante in {percorso_excel}")
    dati_impostazioni = leggi_foglio_impostazioni(wb["Impostazioni"])

    # Foglio Intenzioni (opzionale, potrebbe non esserci)
    intenzioni: list[Intenzione] = []
    if "Intenzioni" in wb.sheetnames:
        intenzioni = leggi_foglio_intenzioni(wb["Intenzioni"])

    # Foglio Matrimoni
    matrimoni: list[Matrimonio] = []
    if "Matrimoni" in wb.sheetnames:
        matrimoni = leggi_foglio_matrimoni(wb["Matrimoni"])

    # Foglio Note
    note: list[dict] = []
    if "Note" in wb.sheetnames:
        note = leggi_foglio_note(wb["Note"])

    # Chiudi il workbook (buona pratica)
    wb.close()

    # Combina i dati Excel
    dati_excel = {
        **dati_impostazioni,
        "intenzioni": intenzioni,
        "matrimoni": matrimoni,
        "note": note,
    }

    # ------------------------------------------------------------------
    # Unione
    # ------------------------------------------------------------------
    return unisci_config(dati_yaml, dati_excel)

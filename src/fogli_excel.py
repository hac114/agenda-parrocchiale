"""
Generatore dei 4 fogli Excel del file config.xlsx.

Contiene le funzioni che creano i singoli fogli:
- Impostazioni: dati generali + tabella periodi
- Intenzioni: registro intenzioni di Messa
- Matrimoni: prenotazioni matrimoni
- Note: annotazioni libere

Questo modulo è usato da `crea_config_template.py`, che orchestra
la creazione del file e gestisce backup e CLI.

Ogni funzione crea un foglio nel Workbook passato come parametro.
"""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from stili import (
    BORDO_SOTTILE,
    FILL_GIALLO,
    FILL_INTESTAZIONE,
    FONT_INTESTAZIONE,
    GIALLO_COMPILABILE,
    stile_compilabile,
    stile_etichetta,
    stile_nota,
    stile_titolo,
)

# ======================================================================
# FOGLIO 1: IMPOSTAZIONI
# ======================================================================


def crea_foglio_impostazioni(wb: Workbook, anno: int, nome_parrocchia: str, citta: str) -> None:
    """Crea il foglio Impostazioni.

    Contiene:
    - Dati generali (anno, nome parrocchia, città) — compilabili
    - Tabella PERIODI: Dal / Al / Feriali / Festivi — compilabili

    Le righe standard sono precompilate; l'utente può modificarle,
    cancellarle o aggiungerne di nuove nelle righe vuote.
    Le righe completamente vuote vengono ignorate in fase di lettura.

    Nota sul formato Dal/Al:
    Le celle "Dal" e "Al" contengono DATE COMPLETE con anno
    (formato gg/mm/aaaa, es. 01/10/2027).
    Anche se i periodi si ripetono ogni anno, l'anno va scritto:
    garantisce validazione forte e permette di usare i filtri Excel.
    Il codice Python leggerà solo mese e giorno (l'anno verrà ignorato
    durante il calcolo, perché i periodi si applicano ogni anno).
    """
    ws = wb.create_sheet("Impostazioni")

    # --- Titolo (unito A1:D1, centrato) ---
    stile_titolo(ws, "A1", "CONFIGURAZIONE AGENDA", larghezza_merge=4)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # ------------------------------------------------------------------
    # SEZIONE 1 — DATI GENERALI
    # ------------------------------------------------------------------
    ws["A3"] = "DATI GENERALI"
    ws["A3"].font = FONT_INTESTAZIONE

    stile_etichetta(ws, "A5", "Anno")
    stile_compilabile(ws, "B5", anno, formato="0")

    stile_etichetta(ws, "A6", "Nome parrocchia")
    stile_compilabile(ws, "B6", nome_parrocchia)

    stile_etichetta(ws, "A7", "Città")
    stile_compilabile(ws, "B7", citta)

    # ------------------------------------------------------------------
    # SEZIONE 2 — TABELLA PERIODI E ORARI
    # ------------------------------------------------------------------
    ws["A9"] = "PERIODI E ORARI"
    ws["A9"].font = FONT_INTESTAZIONE

    stile_nota(
        ws,
        "A10",
        "Compila i periodi che ti servono. Le righe completamente vuote vengono ignorate.\n"
        "Date in formato gg/mm/aaaa (es. 01/10/2027). Orari separati da virgola (es. 7:00,10:00,18:00).",
    )

    # Intestazioni della tabella periodi (riga 12)
    riga_intestazioni = 12
    intestazioni_periodi = ["Dal", "Al", "Feriali", "Festivi"]
    for i, intest in enumerate(intestazioni_periodi, start=1):
        cella = ws.cell(row=riga_intestazioni, column=i, value=intest)
        cella.font = FONT_INTESTAZIONE
        cella.fill = FILL_INTESTAZIONE
        cella.border = BORDO_SOTTILE
        cella.alignment = Alignment(horizontal="center", vertical="center")

    # Righe periodi: 5 standard precompilati + 5 vuoti per aggiunte
    # NOTA: le date includono l'anno corrente (parametro `anno`)
    periodi_standard = [
        (
            date(anno, 10, 1),
            date(anno + 1, 4, 30),
            "7:00,10:00,18:00",
            "8:30,10:00,11:30,18:00",
        ),
        (
            date(anno, 6, 1),
            date(anno, 6, 30),
            "7:00,10:00,19:00",
            "8:30,10:00,11:30,19:00",
        ),
        (
            date(anno, 7, 1),
            date(anno, 8, 31),
            "7:00,10:00,19:00",
            "8:30,10:00,11:30,21:00",
        ),
        (
            date(anno, 9, 1),
            date(anno, 9, 30),
            "7:00,10:00,19:00",
            "8:30,10:00,11:30,19:00",
        ),
        (
            date(anno, 5, 1),
            date(anno, 5, 31),
            "6:15,7:00,8:30,10:00,11:30,17:30,18:30",
            "6:15,7:00,8:30,10:00,11:30,17:30,18:30",
        ),
    ]

    righe_vuote = 5

    riga_inizio_periodi = riga_intestazioni + 1
    riga_fine_periodi = riga_inizio_periodi + len(periodi_standard) + righe_vuote - 1

    for i, (dal, al, feriali, festivi) in enumerate(periodi_standard):
        riga = riga_inizio_periodi + i
        stile_compilabile(ws, f"A{riga}", dal, formato="DD/MM/YYYY")
        stile_compilabile(ws, f"B{riga}", al, formato="DD/MM/YYYY")
        stile_compilabile(ws, f"C{riga}", feriali)
        stile_compilabile(ws, f"D{riga}", festivi)

    for i in range(righe_vuote):
        riga = riga_inizio_periodi + len(periodi_standard) + i
        stile_compilabile(ws, f"A{riga}", formato="DD/MM/YYYY")
        stile_compilabile(ws, f"B{riga}", formato="DD/MM/YYYY")
        stile_compilabile(ws, f"C{riga}")
        stile_compilabile(ws, f"D{riga}")

    # ------------------------------------------------------------------
    # LARGHEZZE COLONNE
    # ------------------------------------------------------------------
    ws.column_dimensions["A"].width = 25  # Dal + etichette
    ws.column_dimensions["B"].width = 14  # Al (data gg/mm/aaaa)
    ws.column_dimensions["C"].width = 42  # Feriali (7 orari a maggio)
    ws.column_dimensions["D"].width = 42  # Festivi

    # ------------------------------------------------------------------
    # VALIDAZIONE INPUT
    # ------------------------------------------------------------------
    # Colonne Dal/Al: date valide nel range 2020-2100
    dv_data = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Data non valida",
        error="Inserisci una data valida in formato gg/mm/aaaa (es. 01/10/2027).",
    )
    ws.add_data_validation(dv_data)
    dv_data.add(f"A{riga_inizio_periodi}:B{riga_fine_periodi}")

    # ------------------------------------------------------------------
    # PROTEZIONE FOGLIO
    # ------------------------------------------------------------------
    for row in ws.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    for row in ws.iter_rows():
        for cell in row:
            if cell.fill.fgColor.rgb in ("00" + GIALLO_COMPILABILE, GIALLO_COMPILABILE):
                cell.protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.enable()
    ws.sheet_view.showGridLines = False


# ======================================================================
# FOGLIO 2: INTENZIONI
# ======================================================================


def crea_foglio_intenzioni(wb: Workbook, righe: int = 400) -> None:
    """Crea il foglio Intenzioni.

    Colonne:
    - A: N. (numerazione automatica tramite formula, non compilabile)
    - B: Data consegna (data, compilabile)
    - C: Intenzione (testo libero, compilabile)
    - D: Offerta in € (numero positivo, compilabile)
    - E: Data applicazione (data, compilabile)
    - F: Note (testo libero, compilabile)

    La numerazione in colonna 'A' appare solo se la colonna 'B' contiene una data.

    Note:
    - La validazione sulle date (B, E) è limitata al range 2020-2100.
      Il controllo "troppo restrittivo" che avevamo in precedenza è stato
      mantenuto, ma va verificato su LibreOffice.
    - La validazione sull'offerta (D) garantisce un importo >= 0.
    - La larghezza della colonna A è 8 (non 6) per evitare l'effetto "###".
    """
    ws = wb.create_sheet("Intenzioni")

    # ------------------------------------------------------------------
    # INTESTAZIONI
    # ------------------------------------------------------------------
    intestazioni = ["N.", "Data consegna", "Intenzione", "Offerta (€)", "Data applicazione", "Note"]

    for i, intest in enumerate(intestazioni, start=1):
        cella = ws.cell(row=1, column=i, value=intest)
        cella.font = FONT_INTESTAZIONE
        cella.fill = FILL_INTESTAZIONE
        cella.border = BORDO_SOTTILE
        cella.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,  # manda a capo il testo
        )

    # Alza la riga delle intestazioni per far spazio al testo a capo
    ws.row_dimensions[1].height = 30

    # ------------------------------------------------------------------
    # RIGHE DATI
    # ------------------------------------------------------------------
    for r in range(2, righe + 2):
        # Colonna A: formula numerazione automatica (protetta, non compilabile)
        cella_a = ws.cell(row=r, column=1)
        cella_a.value = f'=IF(B{r}="","",ROW()-1)'
        cella_a.border = BORDO_SOTTILE
        cella_a.alignment = Alignment(horizontal="center", vertical="center")
        cella_a.number_format = "0"

        # Colonne B–F: gialle, compilabili
        for c in range(2, 7):
            cella = ws.cell(row=r, column=c)
            cella.fill = FILL_GIALLO
            cella.border = BORDO_SOTTILE
            cella.alignment = Alignment(horizontal="left", vertical="center")
            if c in (2, 5):
                # Date consegna e applicazione
                cella.number_format = "DD/MM/YYYY"
            elif c == 4:
                # Offerta
                cella.number_format = "€ #,##0.00"

    # ------------------------------------------------------------------
    # LARGHEZZE COLONNE
    # ------------------------------------------------------------------
    ws.column_dimensions["A"].width = 8  # N. (aumentata: evita ### )
    ws.column_dimensions["B"].width = 14  # Data consegna
    ws.column_dimensions["C"].width = 50  # Intenzione (testo lungo)
    ws.column_dimensions["D"].width = 12  # Offerta
    ws.column_dimensions["E"].width = 14  # Data applicazione
    ws.column_dimensions["F"].width = 30  # Note

    # ------------------------------------------------------------------
    # VALIDAZIONE INPUT
    # ------------------------------------------------------------------
    # Colonne B ed E: date valide tra 2020 e 2100
    dv_data = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Data non valida",
        error="Inserisci una data valida (gg/mm/aaaa).",
    )
    ws.add_data_validation(dv_data)
    dv_data.add(f"B2:B{righe + 1}")
    dv_data.add(f"E2:E{righe + 1}")

    # Colonna D: offerta numero positivo o zero
    dv_offerta = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Offerta non valida",
        error="Inserisci un importo positivo (es. 10.00).",
    )
    ws.add_data_validation(dv_offerta)
    dv_offerta.add(f"D2:D{righe + 1}")

    # ------------------------------------------------------------------
    # PROTEZIONE FOGLIO
    # ------------------------------------------------------------------
    # Blocca tutto, poi sblocca solo le celle gialle (B–F)
    for row in ws.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    for r in range(2, righe + 2):
        for c in range(2, 7):
            ws.cell(row=r, column=c).protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.enable()

    # Blocca la prima riga (intestazioni) durante lo scorrimento
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


# ======================================================================
# FOGLIO 3: MATRIMONI
# ======================================================================


def crea_foglio_matrimoni(wb: Workbook, righe: int = 100) -> None:
    """Crea il foglio Matrimoni.

    Contiene le prenotazioni dei matrimoni per l'anno successivo.
    Colonne:
    - A: Data (data, compilabile)
    - B: Ora (ora nativa HH:MM, compilabile)
    - C: Nome sposi (testo libero, compilabile)
    - D: Contatti (testo libero, compilabile)
    - E: Note (testo libero, compilabile)

    Nota: nell'agenda stampata i matrimoni compaiono in fondo (ultima pagina)
    con data, ora e nome sposi, come richiesto dal santuario.

    Note tecniche:
    - Nessuna validazione Excel sulle date/ore: il formato cella
      (DD/MM/YYYY e HH:MM) è sufficiente per la visualizzazione.
      La validazione forte sarà fatta dal parser Python in lettura.
    - Le intestazioni hanno wrap_text attivo per andare a capo se lunghe.
    """
    ws = wb.create_sheet("Matrimoni")

    # ------------------------------------------------------------------
    # INTESTAZIONI
    # ------------------------------------------------------------------
    intestazioni = ["Data", "Ora", "Nome sposi", "Contatti", "Note"]

    for i, intest in enumerate(intestazioni, start=1):
        cella = ws.cell(row=1, column=i, value=intest)
        cella.font = FONT_INTESTAZIONE
        cella.fill = FILL_INTESTAZIONE
        cella.border = BORDO_SOTTILE
        cella.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Alza la riga delle intestazioni per far spazio al testo a capo
    ws.row_dimensions[1].height = 30

    # ------------------------------------------------------------------
    # RIGHE DATI
    # ------------------------------------------------------------------
    for r in range(2, righe + 2):
        for c in range(1, 6):
            cella = ws.cell(row=r, column=c)
            cella.fill = FILL_GIALLO
            cella.border = BORDO_SOTTILE
            cella.alignment = Alignment(horizontal="left", vertical="center")
            if c == 1:
                # Colonna Data
                cella.number_format = "DD/MM/YYYY"
            elif c == 2:
                # Colonna Ora → formato ora nativo
                cella.number_format = "HH:MM"

    # ------------------------------------------------------------------
    # LARGHEZZE COLONNE
    # ------------------------------------------------------------------
    ws.column_dimensions["A"].width = 14  # Data
    ws.column_dimensions["B"].width = 10  # Ora
    ws.column_dimensions["C"].width = 35  # Nome sposi
    ws.column_dimensions["D"].width = 25  # Contatti
    ws.column_dimensions["E"].width = 35  # Note

    # ------------------------------------------------------------------
    # PROTEZIONE FOGLIO
    # ------------------------------------------------------------------
    for row in ws.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    for r in range(2, righe + 2):
        for c in range(1, 6):
            ws.cell(row=r, column=c).protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.enable()

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


# ======================================================================
# FOGLIO 4: NOTE
# ======================================================================


def crea_foglio_note(wb: Workbook, righe: int = 30) -> None:
    """Crea il foglio Note.

    Annotazioni libere per periodi o giorni speciali.
    Colonne:
    - A: Dal (data inizio, compilabile)
    - B: Al (data fine, compilabile)
    - C: Nota (testo lungo, compilabile)

    Comportamento:
    - Se Al è vuoto e Dal è compilato, il sistema Python considera
      il periodo come un giorno singolo (Al = Dal).
    - L'utente NON deve compilare Al per note di un solo giorno.

    Uso tipico:
    - Dal=14/03/2027, Al=16/03/2027, Nota="Triduo San Salvatore"
    - Dal=01/05/2027, Al=31/05/2027, Nota="Mese mariano, programma da definire"
    - Dal=30/05/2027, (Al vuoto), Nota="Corpus Domini, vietato pomeriggio"

    Note tecniche:
    - Le intestazioni hanno wrap_text attivo.
    - Formato date su Dal/Al, formato testo su Nota.
    - Validazione data su Dal e Al (range 2020-2100).
    """
    ws = wb.create_sheet("Note")

    # ------------------------------------------------------------------
    # INTESTAZIONI
    # ------------------------------------------------------------------
    intestazioni = ["Dal", "Al", "Nota"]

    for i, intest in enumerate(intestazioni, start=1):
        cella = ws.cell(row=1, column=i, value=intest)
        cella.font = FONT_INTESTAZIONE
        cella.fill = FILL_INTESTAZIONE
        cella.border = BORDO_SOTTILE
        cella.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Alza la riga delle intestazioni per far spazio al testo a capo
    ws.row_dimensions[1].height = 30

    # ------------------------------------------------------------------
    # RIGHE DATI
    # ------------------------------------------------------------------
    for r in range(2, righe + 2):
        for c in range(1, 4):
            cella = ws.cell(row=r, column=c)
            cella.fill = FILL_GIALLO
            cella.border = BORDO_SOTTILE
            if c in (1, 2):
                # Colonne Dal/Al → data
                cella.number_format = "DD/MM/YYYY"
                cella.alignment = Alignment(horizontal="left", vertical="center")
            else:
                # Colonna Nota → testo con a capo automatico
                cella.alignment = Alignment(
                    horizontal="left",
                    vertical="top",
                    wrap_text=True,
                )

    # ------------------------------------------------------------------
    # LARGHEZZE COLONNE
    # ------------------------------------------------------------------
    ws.column_dimensions["A"].width = 14  # Dal
    ws.column_dimensions["B"].width = 14  # Al
    ws.column_dimensions["C"].width = 80  # Nota

    # ------------------------------------------------------------------
    # VALIDAZIONE INPUT
    # ------------------------------------------------------------------
    # Colonne Dal/Al: date valide tra 2020 e 2100
    dv_data = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Data non valida",
        error="Inserisci una data valida (gg/mm/aaaa).",
    )
    ws.add_data_validation(dv_data)
    dv_data.add(f"A2:B{righe + 1}")

    # ------------------------------------------------------------------
    # PROTEZIONE FOGLIO
    # ------------------------------------------------------------------
    # Blocca tutto, poi sblocca solo le celle gialle (A–C)
    for row in ws.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    for r in range(2, righe + 2):
        for c in range(1, 4):
            ws.cell(row=r, column=c).protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.enable()

    # Blocca la prima riga (intestazioni) durante lo scorrimento
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

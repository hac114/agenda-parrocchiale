"""Test per il modulo lettura_excel."""

from __future__ import annotations

from datetime import date
from typing import cast

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from lettura_excel import (
    leggi_foglio_impostazioni,
    leggi_foglio_intenzioni,
    leggi_foglio_matrimoni,
    leggi_foglio_note,
    parse_orari,
)

# ======================================================================
# HELPER
# ======================================================================


def crea_foglio_vuoto() -> Worksheet:
    """Crea un foglio Excel vuoto per i test.

    Il cast è necessario perché openpyxl tipizza `wb.active` come
    `Worksheet | None`, anche se per un Workbook appena creato è
    sempre valorizzato.
    """
    wb = Workbook()
    return cast(Worksheet, wb.active)


# ======================================================================
# TEST — parse_orari
# ======================================================================


def test_parse_orari_virgole() -> None:
    """Orari separati da virgola."""
    assert parse_orari("7:00,10:00,18:00") == ["7:00", "10:00", "18:00"]


def test_parse_orari_punto_virgola() -> None:
    """Orari separati da punto e virgola."""
    assert parse_orari("7:00;10:00") == ["7:00", "10:00"]


def test_parse_orari_spazi() -> None:
    """Orari separati da spazi."""
    assert parse_orari("7:00 10:00 18:00") == ["7:00", "10:00", "18:00"]


def test_parse_orari_vuoto() -> None:
    """Stringa vuota → lista vuota."""
    assert parse_orari("") == []


def test_parse_orari_valori_invalidi() -> None:
    """Valori non validi vengono ignorati."""
    assert parse_orari("7:00,pippo,18:00") == ["7:00", "18:00"]
    assert parse_orari("7.00,10:00") == ["10:00"]


def test_parse_orari_misto() -> None:
    """Separatori misti."""
    assert parse_orari("7:00, 10:00; 18:00") == ["7:00", "10:00", "18:00"]


# ======================================================================
# FIXTURE — Foglio Impostazioni
# ======================================================================


@pytest.fixture
def foglio_impostazioni_valido() -> Worksheet:
    """Crea un foglio Impostazioni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Impostazioni"

    ws["B5"] = 2027
    ws["B6"] = "San Pietro in Silki"
    ws["B7"] = "Sassari"

    ws.cell(row=13, column=1, value=date(2027, 10, 1))
    ws.cell(row=13, column=2, value=date(2028, 4, 30))
    ws.cell(row=13, column=3, value="7:00,10:00,18:00")
    ws.cell(row=13, column=4, value="8:30,10:00,11:30,18:00")

    ws.cell(row=14, column=1, value=date(2027, 5, 1))
    ws.cell(row=14, column=2, value=date(2027, 5, 31))
    ws.cell(row=14, column=3, value="6:15,7:00,8:30,10:00,11:30,17:30,18:30")
    ws.cell(row=14, column=4, value="6:15,7:00,8:30,10:00,11:30,17:30,18:30")

    return ws


# ======================================================================
# TEST — leggi_foglio_impostazioni
# ======================================================================


def test_leggi_impostazioni_valido(foglio_impostazioni_valido: Worksheet) -> None:
    """Un foglio valido viene letto correttamente."""
    dati = leggi_foglio_impostazioni(foglio_impostazioni_valido)
    assert dati["anno"] == 2027
    assert dati["nome_parrocchia"] == "San Pietro in Silki"
    assert dati["citta"] == "Sassari"
    assert len(dati["periodi"]) == 2
    assert dati["periodi"][0]["dal"] == date(2027, 10, 1)
    assert dati["periodi"][0]["orari_feriali"] == ["7:00", "10:00", "18:00"]
    assert dati["periodi"][1]["orari_festivi"] == [
        "6:15",
        "7:00",
        "8:30",
        "10:00",
        "11:30",
        "17:30",
        "18:30",
    ]


def test_leggi_impostazioni_anno_mancante() -> None:
    """Se manca l'anno → ValueError."""
    ws = crea_foglio_vuoto()
    ws["B5"] = None
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    with pytest.raises(ValueError, match="Anno"):
        leggi_foglio_impostazioni(ws)


def test_leggi_impostazioni_nome_mancante() -> None:
    """Se manca il nome parrocchia → ValueError."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = None
    ws["B7"] = "Test"

    with pytest.raises(ValueError, match="Nome parrocchia"):
        leggi_foglio_impostazioni(ws)


def test_leggi_impostazioni_riga_parziale() -> None:
    """Una riga con solo 'Dal' viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    ws.cell(row=13, column=1, value=date(2027, 10, 1))
    ws.cell(row=13, column=3, value="7:00")

    dati = leggi_foglio_impostazioni(ws)
    assert len(dati["periodi"]) == 0


def test_leggi_impostazioni_riga_non_data() -> None:
    """Una riga con valore non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    ws.cell(row=13, column=1, value="01/10/2027")
    ws.cell(row=13, column=2, value="30/04/2028")

    dati = leggi_foglio_impostazioni(ws)
    assert len(dati["periodi"]) == 0


# ======================================================================
# FIXTURE — Foglio Intenzioni
# ======================================================================


@pytest.fixture
def foglio_intenzioni_valido() -> Worksheet:
    """Crea un foglio Intenzioni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value="Per la pace nel mondo")
    ws.cell(row=2, column=4, value=10.0)
    ws.cell(row=2, column=5, value=date(2027, 1, 20))
    ws.cell(row=2, column=6, value="Richiesta da Mario")

    ws.cell(row=3, column=2, value=date(2027, 2, 1))
    ws.cell(row=3, column=3, value="Per i defunti della famiglia Rossi")
    ws.cell(row=3, column=4, value=20.0)

    ws.cell(row=5, column=2, value=date(2027, 3, 10))
    ws.cell(row=5, column=3, value="Per la salute di Anna")

    return ws


# ======================================================================
# TEST — leggi_foglio_intenzioni
# ======================================================================


def test_leggi_intenzioni_valido(foglio_intenzioni_valido: Worksheet) -> None:
    """Un foglio valido restituisce 3 intenzioni."""
    intenzioni = leggi_foglio_intenzioni(foglio_intenzioni_valido)
    assert len(intenzioni) == 3

    i1 = intenzioni[0]
    assert i1.numero == 1
    assert i1.data_consegna == date(2027, 1, 15)
    assert i1.testo == "Per la pace nel mondo"
    assert i1.offerta == 10.0
    assert i1.data_applicazione == date(2027, 1, 20)
    assert i1.note == "Richiesta da Mario"

    i2 = intenzioni[1]
    assert i2.numero == 2
    assert i2.data_applicazione is None
    assert i2.offerta == 20.0

    i3 = intenzioni[2]
    assert i3.numero == 3
    assert i3.offerta is None


def test_leggi_intenzioni_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    assert leggi_foglio_intenzioni(ws) == []


def test_leggi_intenzioni_riga_senza_testo() -> None:
    """Una riga con Data ma senza Intenzione viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value=None)

    ws.cell(row=3, column=2, value=date(2027, 2, 1))
    ws.cell(row=3, column=3, value="Intenzione valida")

    intenzioni = leggi_foglio_intenzioni(ws)
    assert len(intenzioni) == 1
    assert intenzioni[0].testo == "Intenzione valida"


def test_leggi_intenzioni_data_non_valida() -> None:
    """Una riga con Data non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    ws.cell(row=2, column=2, value="15/01/2027")
    ws.cell(row=2, column=3, value="Testo")

    assert leggi_foglio_intenzioni(ws) == []


def test_leggi_intenzioni_offerta_non_numerica() -> None:
    """Un'offerta non numerica viene impostata a None con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value="Testo valido")
    ws.cell(row=2, column=4, value="dieci euro")

    intenzioni = leggi_foglio_intenzioni(ws)
    assert len(intenzioni) == 1
    assert intenzioni[0].offerta is None


# ======================================================================
# FIXTURE — Foglio Matrimoni
# ======================================================================


@pytest.fixture
def foglio_matrimoni_valido() -> Worksheet:
    """Crea un foglio Matrimoni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value="Maria Rossi e Luca Bianchi")
    ws.cell(row=2, column=4, value="333-1234567")
    ws.cell(row=2, column=5, value="Ricevimento in parrocchia")

    ws.cell(row=3, column=1, value=date(2028, 7, 15))
    ws.cell(row=3, column=2, value="10:00")
    ws.cell(row=3, column=3, value="Anna Verdi e Marco Neri")

    return ws


# ======================================================================
# TEST — leggi_foglio_matrimoni
# ======================================================================


def test_leggi_matrimoni_valido(foglio_matrimoni_valido: Worksheet) -> None:
    """Un foglio valido restituisce 2 matrimoni."""
    matrimoni = leggi_foglio_matrimoni(foglio_matrimoni_valido)
    assert len(matrimoni) == 2

    m1 = matrimoni[0]
    assert m1.data == date(2028, 6, 10)
    assert m1.ora == "15:30"
    assert m1.nome_sposi == "Maria Rossi e Luca Bianchi"
    assert m1.contatti == "333-1234567"
    assert m1.note == "Ricevimento in parrocchia"

    m2 = matrimoni[1]
    assert m2.data == date(2028, 7, 15)
    assert m2.contatti == ""
    assert m2.note == ""


def test_leggi_matrimoni_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_senza_ora() -> None:
    """Una riga senza Ora viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value=None)
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_ora_non_valida() -> None:
    """Una riga con Ora malformata viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15.30")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_senza_nome() -> None:
    """Una riga senza Nome sposi viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value=None)

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_data_non_data() -> None:
    """Una riga con Data non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value="10/06/2028")
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_ora_nativa() -> None:
    """Un'ora nativa (datetime.time) viene convertita in stringa HH:MM."""
    from datetime import time

    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value=time(15, 30))
    ws.cell(row=2, column=3, value="Test")

    matrimoni = leggi_foglio_matrimoni(ws)
    assert len(matrimoni) == 1
    assert matrimoni[0].ora == "15:30"


# ======================================================================
# FIXTURE — Foglio Note
# ======================================================================


@pytest.fixture
def foglio_note_valido() -> Worksheet:
    """Crea un foglio Note valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value=date(2027, 3, 16))
    ws.cell(row=2, column=3, value="Triduo San Salvatore")

    ws.cell(row=3, column=1, value=date(2027, 5, 30))
    ws.cell(row=3, column=3, value="Corpus Domini, vietato pomeriggio")

    ws.cell(row=5, column=1, value=date(2027, 5, 1))
    ws.cell(row=5, column=2, value=date(2027, 5, 31))
    ws.cell(row=5, column=3, value="Mese mariano, programma da definire")

    return ws


# ======================================================================
# TEST — leggi_foglio_note
# ======================================================================


def test_leggi_note_valido(foglio_note_valido: Worksheet) -> None:
    """Un foglio valido restituisce 3 note."""
    note = leggi_foglio_note(foglio_note_valido)
    assert len(note) == 3

    n1 = note[0]
    assert n1["dal"] == date(2027, 3, 14)
    assert n1["al"] == date(2027, 3, 16)
    assert n1["nota"] == "Triduo San Salvatore"

    n2 = note[1]
    assert n2["dal"] == date(2027, 5, 30)
    assert n2["al"] == date(2027, 5, 30)

    n3 = note[2]
    assert n3["dal"] == date(2027, 5, 1)
    assert n3["al"] == date(2027, 5, 31)


def test_leggi_note_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    assert leggi_foglio_note(ws) == []


def test_leggi_note_senza_nota() -> None:
    """Una riga con Dal ma senza Nota viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=3, value=None)

    assert leggi_foglio_note(ws) == []


def test_leggi_note_dal_non_data() -> None:
    """Una riga con Dal non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value="14/03/2027")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_non_data() -> None:
    """Una riga con Al non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value="16/03/2027")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_prima_di_dal() -> None:
    """Una riga con Al < Dal viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 16))
    ws.cell(row=2, column=2, value=date(2027, 3, 14))
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_uguale_dal() -> None:
    """Una nota con Al = Dal è valida."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value=date(2027, 3, 14))
    ws.cell(row=2, column=3, value="Test")

    note = leggi_foglio_note(ws)
    assert len(note) == 1
    assert note[0]["dal"] == note[0]["al"]

"""Test per il modulo fogli_excel."""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook

from fogli_excel import (
    crea_foglio_impostazioni,
    crea_foglio_intenzioni,
    crea_foglio_matrimoni,
    crea_foglio_note,
)

# ======================================================================
# FIXTURE
# ======================================================================


@pytest.fixture
def wb_vuoto() -> Workbook:
    """Workbook vuoto."""
    return Workbook()


# ======================================================================
# TEST — crea_foglio_impostazioni
# ======================================================================


def test_foglio_impostazioni_creato(wb_vuoto: Workbook) -> None:
    """Il foglio Impostazioni viene creato."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test Parrocchia", "Roma")
    assert "Impostazioni" in wb_vuoto.sheetnames


def test_foglio_impostazioni_dati_generali(wb_vuoto: Workbook) -> None:
    """I dati generali sono scritti nelle celle corrette."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test Parrocchia", "Roma")
    ws = wb_vuoto["Impostazioni"]

    assert ws["B5"].value == 2027
    assert ws["B6"].value == "Test Parrocchia"
    assert ws["B7"].value == "Roma"


def test_foglio_impostazioni_titolo(wb_vuoto: Workbook) -> None:
    """Il titolo è scritto in A1."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test", "Test")
    ws = wb_vuoto["Impostazioni"]

    assert ws["A1"].value == "CONFIGURAZIONE AGENDA"


def test_foglio_impostazioni_periodi_precompilati(wb_vuoto: Workbook) -> None:
    """Le 5 righe standard dei periodi sono precompilate."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test", "Test")
    ws = wb_vuoto["Impostazioni"]

    # Prima riga periodi (13): inverno
    assert ws["A13"].value == date(2027, 10, 1)
    assert ws["B13"].value == date(2028, 4, 30)

    # Seconda riga (14): giugno
    assert ws["A14"].value == date(2027, 6, 1)
    assert ws["B14"].value == date(2027, 6, 30)

    # Terza riga (15): luglio/agosto
    assert ws["A15"].value == date(2027, 7, 1)

    # Quarta riga (16): settembre
    assert ws["A16"].value == date(2027, 9, 1)

    # Quinta riga (17): maggio
    assert ws["A17"].value == date(2027, 5, 1)
    assert ws["B17"].value == date(2027, 5, 31)


def test_foglio_impostazioni_celle_compilabili_gialle(wb_vuoto: Workbook) -> None:
    """Le celle compilabili (Anno, Nome, Città) sono gialle."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test", "Test")
    ws = wb_vuoto["Impostazioni"]

    # B5 (Anno), B6 (Nome), B7 (Città) devono avere riempimento giallo
    assert ws["B5"].fill.fgColor.rgb in ("00FFF9C4", "FFF9C4")
    assert ws["B6"].fill.fgColor.rgb in ("00FFF9C4", "FFF9C4")
    assert ws["B7"].fill.fgColor.rgb in ("00FFF9C4", "FFF9C4")


def test_foglio_impostazioni_protezione_attiva(wb_vuoto: Workbook) -> None:
    """Il foglio è protetto."""
    crea_foglio_impostazioni(wb_vuoto, 2027, "Test", "Test")
    ws = wb_vuoto["Impostazioni"]

    assert ws.protection.sheet is True


# ======================================================================
# TEST — crea_foglio_intenzioni
# ======================================================================


def test_foglio_intenzioni_creato(wb_vuoto: Workbook) -> None:
    """Il foglio Intenzioni viene creato."""
    crea_foglio_intenzioni(wb_vuoto)
    assert "Intenzioni" in wb_vuoto.sheetnames


def test_foglio_intenzioni_intestazioni(wb_vuoto: Workbook) -> None:
    """Le intestazioni sono scritte correttamente."""
    crea_foglio_intenzioni(wb_vuoto)
    ws = wb_vuoto["Intenzioni"]

    assert ws["A1"].value == "N."
    assert ws["B1"].value == "Data consegna"
    assert ws["C1"].value == "Intenzione"
    assert ws["D1"].value == "Offerta (€)"
    assert ws["E1"].value == "Data applicazione"
    assert ws["F1"].value == "Note"


def test_foglio_intenzioni_formula_numerazione(wb_vuoto: Workbook) -> None:
    """La colonna A ha la formula di numerazione."""
    crea_foglio_intenzioni(wb_vuoto, righe=10)
    ws = wb_vuoto["Intenzioni"]

    # Riga 2
    assert ws["A2"].value == '=IF(B2="","",ROW()-1)'
    # Riga 11 (ultima riga)
    assert ws["A11"].value == '=IF(B11="","",ROW()-1)'


def test_foglio_intenzioni_celle_compilabili(wb_vuoto: Workbook) -> None:
    """Le colonne B–F sono gialle."""
    crea_foglio_intenzioni(wb_vuoto, righe=10)
    ws = wb_vuoto["Intenzioni"]

    for col in "BCDEF":
        cella = ws[f"{col}2"]
        assert cella.fill.fgColor.rgb in ("00FFF9C4", "FFF9C4"), f"Colonna {col} non gialla"


def test_foglio_intenzioni_protezione(wb_vuoto: Workbook) -> None:
    """Il foglio è protetto."""
    crea_foglio_intenzioni(wb_vuoto)
    ws = wb_vuoto["Intenzioni"]

    assert ws.protection.sheet is True


def test_foglio_intenzioni_freeze_panes(wb_vuoto: Workbook) -> None:
    """Le intestazioni sono bloccate (freeze panes)."""
    crea_foglio_intenzioni(wb_vuoto)
    ws = wb_vuoto["Intenzioni"]

    assert ws.freeze_panes == "A2"


# ======================================================================
# TEST — crea_foglio_matrimoni
# ======================================================================


def test_foglio_matrimoni_creato(wb_vuoto: Workbook) -> None:
    """Il foglio Matrimoni viene creato."""
    crea_foglio_matrimoni(wb_vuoto)
    assert "Matrimoni" in wb_vuoto.sheetnames


def test_foglio_matrimoni_intestazioni(wb_vuoto: Workbook) -> None:
    """Le intestazioni sono scritte correttamente."""
    crea_foglio_matrimoni(wb_vuoto)
    ws = wb_vuoto["Matrimoni"]

    assert ws["A1"].value == "Data"
    assert ws["B1"].value == "Ora"
    assert ws["C1"].value == "Nome sposi"
    assert ws["D1"].value == "Contatti"
    assert ws["E1"].value == "Note"


def test_foglio_matrimoni_formato_date(wb_vuoto: Workbook) -> None:
    """La colonna A ha formato data."""
    crea_foglio_matrimoni(wb_vuoto, righe=5)
    ws = wb_vuoto["Matrimoni"]

    assert ws["A2"].number_format == "DD/MM/YYYY"


def test_foglio_matrimoni_formato_ora(wb_vuoto: Workbook) -> None:
    """La colonna B ha formato ora."""
    crea_foglio_matrimoni(wb_vuoto, righe=5)
    ws = wb_vuoto["Matrimoni"]

    assert ws["B2"].number_format == "HH:MM"


# ======================================================================
# TEST — crea_foglio_note
# ======================================================================


def test_foglio_note_creato(wb_vuoto: Workbook) -> None:
    """Il foglio Note viene creato."""
    crea_foglio_note(wb_vuoto)
    assert "Note" in wb_vuoto.sheetnames


def test_foglio_note_intestazioni(wb_vuoto: Workbook) -> None:
    """Le intestazioni sono scritte correttamente."""
    crea_foglio_note(wb_vuoto)
    ws = wb_vuoto["Note"]

    assert ws["A1"].value == "Dal"
    assert ws["B1"].value == "Al"
    assert ws["C1"].value == "Nota"


def test_foglio_note_formato_date(wb_vuoto: Workbook) -> None:
    """Le colonne A e B hanno formato data."""
    crea_foglio_note(wb_vuoto, righe=5)
    ws = wb_vuoto["Note"]

    assert ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["B2"].number_format == "DD/MM/YYYY"


def test_foglio_note_celle_compilabili(wb_vuoto: Workbook) -> None:
    """Le colonne A, B, C sono gialle."""
    crea_foglio_note(wb_vuoto, righe=5)
    ws = wb_vuoto["Note"]

    for col in "ABC":
        cella = ws[f"{col}2"]
        assert cella.fill.fgColor.rgb in ("00FFF9C4", "FFF9C4")

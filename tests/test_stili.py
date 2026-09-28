"""Test per il modulo stili."""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from stili import (
    FILL_GIALLO,
    GIALLO_COMPILABILE,
    auto_adatta_colonne,
    proteggi_foglio,
    stile_compilabile,
    stile_etichetta,
    stile_nota,
    stile_titolo,
)

# ======================================================================
# FIXTURE
# ======================================================================


@pytest.fixture
def ws() -> Worksheet:
    """Foglio di lavoro vuoto."""
    wb = Workbook()
    return wb.active  # type: ignore[return-value]


# ======================================================================
# TEST — stile_titolo
# ======================================================================


def test_stile_titolo_scrive_testo(ws: Worksheet) -> None:
    """Il titolo è scritto nella cella."""
    stile_titolo(ws, "A1", "Titolo Test")
    assert ws["A1"].value == "Titolo Test"


def test_stile_titolo_unisce_celle(ws: Worksheet) -> None:
    """Il titolo unisce le celle specificate."""
    stile_titolo(ws, "A1", "Titolo", larghezza_merge=3)
    # Verifica che sia stato unito A1:C1
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert "A1:C1" in merged_ranges


def test_stile_titolo_font_grassetto(ws: Worksheet) -> None:
    """Il titolo ha font grassetto."""
    stile_titolo(ws, "A1", "Titolo")
    assert ws["A1"].font.bold is True


def test_stile_titolo_fill(ws: Worksheet) -> None:
    """Il titolo ha fill colorato."""
    stile_titolo(ws, "A1", "Titolo")
    # Il fill è blu scuro
    assert ws["A1"].fill.fgColor.rgb in ("001F4E79", "1F4E79")


# ======================================================================
# TEST — stile_etichetta
# ======================================================================


def test_stile_etichetta_scrive_testo(ws: Worksheet) -> None:
    """L'etichetta scrive il testo."""
    stile_etichetta(ws, "A1", "Anno")
    assert ws["A1"].value == "Anno"


def test_stile_etichetta_font(ws: Worksheet) -> None:
    """L'etichetta ha font grassetto."""
    stile_etichetta(ws, "A1", "Test")
    assert ws["A1"].font.bold is True


def test_stile_etichetta_bordo(ws: Worksheet) -> None:
    """L'etichetta ha bordo sottile."""
    stile_etichetta(ws, "A1", "Test")
    assert ws["A1"].border.left.style == "thin"


# ======================================================================
# TEST — stile_compilabile
# ======================================================================


def test_stile_compilabile_senza_valore(ws: Worksheet) -> None:
    """La cella compilabile senza valore è solo gialla."""
    stile_compilabile(ws, "A1")
    assert ws["A1"].value is None
    assert ws["A1"].fill.fgColor.rgb in ("00FFF9C4", GIALLO_COMPILABILE)


def test_stile_compilabile_con_valore(ws: Worksheet) -> None:
    """La cella compilabile scrive il valore."""
    stile_compilabile(ws, "A1", "Testo")
    assert ws["A1"].value == "Testo"


def test_stile_compilabile_con_formato(ws: Worksheet) -> None:
    """La cella compilabile applica il formato."""
    stile_compilabile(ws, "A1", 2027, formato="0")
    assert ws["A1"].number_format == "0"


def test_stile_compilabile_fill_giallo(ws: Worksheet) -> None:
    """La cella compilabile è gialla."""
    stile_compilabile(ws, "A1")
    assert ws["A1"].fill == FILL_GIALLO


# ======================================================================
# TEST — stile_nota
# ======================================================================


def test_stile_nota_scrive_testo(ws: Worksheet) -> None:
    """La nota scrive il testo."""
    stile_nota(ws, "A1", "Nota informativa")
    assert ws["A1"].value == "Nota informativa"


def test_stile_nota_font_corsivo(ws: Worksheet) -> None:
    """La nota è in corsivo."""
    stile_nota(ws, "A1", "Test")
    assert ws["A1"].font.italic is True


def test_stile_nota_wrap_text(ws: Worksheet) -> None:
    """La nota ha wrap text attivo."""
    stile_nota(ws, "A1", "Testo lungo da andare a capo")
    assert ws["A1"].alignment.wrap_text is True


# ======================================================================
# TEST — auto_adatta_colonne
# ======================================================================


def test_auto_adatta_colonne_base(ws: Worksheet) -> None:
    """Le colonne vengono adattate al contenuto."""
    ws["A1"] = "Testo corto"
    auto_adatta_colonne(ws)
    assert ws.column_dimensions["A"].width >= 10


def test_auto_adatta_colonne_testo_lungo(ws: Worksheet) -> None:
    """Un testo lungo allarga di più la colonna."""
    ws["A1"] = "Testo molto lungo che occupa molto spazio"
    auto_adatta_colonne(ws)
    assert ws.column_dimensions["A"].width >= 30


def test_auto_adatta_colonne_limite_massimo(ws: Worksheet) -> None:
    """La larghezza ha un limite massimo (60)."""
    ws["A1"] = "x" * 200  # 200 caratteri
    auto_adatta_colonne(ws)
    assert ws.column_dimensions["A"].width <= 60


def test_auto_adatta_colonne_limite_minimo(ws: Worksheet) -> None:
    """La larghezza ha un limite minimo (8)."""
    ws["A1"] = "a"
    auto_adatta_colonne(ws)
    assert ws.column_dimensions["A"].width >= 8


def test_auto_adatta_colonne_ignora_vuote(ws: Worksheet) -> None:
    """Le celle vuote non influenzano la larghezza."""
    ws["A1"] = "Test"
    ws["A2"] = None
    auto_adatta_colonne(ws)
    assert ws.column_dimensions["A"].width >= 8


# ======================================================================
# TEST — proteggi_foglio
# ======================================================================


def test_proteggi_foglio_senza_sblocco(ws: Worksheet) -> None:
    """Il foglio è protetto con tutte le celle bloccate."""
    proteggi_foglio(ws)
    assert ws.protection.sheet is True


def test_proteggi_foglio_blocca_tutto(ws: Worksheet) -> None:
    """Senza celle sbloccate, tutte le celle sono locked."""
    ws["A1"] = "test"
    ws["B2"] = "test"
    proteggi_foglio(ws)
    # Le celle sono locked
    assert ws["A1"].protection.locked is True
    assert ws["B2"].protection.locked is True


def test_proteggi_foglio_sblocca_intervallo(ws: Worksheet) -> None:
    """Le celle nell'intervallo sbloccato sono unlocked."""
    ws["A1"] = "locked"
    ws["B2"] = "unlocked"
    ws["C3"] = "unlocked"
    proteggi_foglio(ws, celle_sbloccate=["B2:C3"])

    # A1 resta locked
    assert ws["A1"].protection.locked is True
    # B2 e C3 sono sbloccati
    assert ws["B2"].protection.locked is False
    assert ws["C3"].protection.locked is False

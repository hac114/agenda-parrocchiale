"""Test per il modulo lettura_config."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import cast

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from lettura_config import (
    carica_config,
    leggi_foglio_impostazioni,
    leggi_foglio_intenzioni,
    leggi_foglio_matrimoni,
    leggi_foglio_note,
    leggi_yaml,
    parse_orari,
    unisci_config,
)

# ======================================================================
# HELPER
# ======================================================================


def crea_foglio_vuoto() -> Worksheet:
    """Crea un foglio Excel vuoto per i test.

    Il cast è necessario perché openpyxl tipizza `wb.active` come
    `Worksheet | None`, anche se per un Workbook appena creato è
    sempre valorizzato.
    """
    wb = Workbook()
    return cast(Worksheet, wb.active)


# ======================================================================
# FIXTURE — YAML
# ======================================================================


@pytest.fixture
def yaml_valido(tmp_path: Path) -> Path:
    """Crea un file YAML valido temporaneo."""
    percorso = tmp_path / "valido.yaml"
    percorso.write_text(
        """
nome: Test
valori:
  - 1
  - 2
  - 3
""",
        encoding="utf-8",
    )
    return percorso


@pytest.fixture
def yaml_vuoto(tmp_path: Path) -> Path:
    """Crea un file YAML vuoto."""
    percorso = tmp_path / "vuoto.yaml"
    percorso.write_text("", encoding="utf-8")
    return percorso


@pytest.fixture
def yaml_malformato(tmp_path: Path) -> Path:
    """Crea un file YAML malformato (tab character dove non ammesso)."""
    percorso = tmp_path / "malformato.yaml"
    percorso.write_text(
        "chiave:\n\tvalore\n",
        encoding="utf-8",
    )
    return percorso


# ======================================================================
# FIXTURE — Foglio Impostazioni
# ======================================================================


@pytest.fixture
def foglio_impostazioni_valido() -> Worksheet:
    """Crea un foglio Impostazioni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Impostazioni"

    # Dati generali
    ws["B5"] = 2027
    ws["B6"] = "San Pietro in Silki"
    ws["B7"] = "Sassari"

    # Periodo 1: inverno
    ws.cell(row=13, column=1, value=date(2027, 10, 1))
    ws.cell(row=13, column=2, value=date(2028, 4, 30))
    ws.cell(row=13, column=3, value="7:00,10:00,18:00")
    ws.cell(row=13, column=4, value="8:30,10:00,11:30,18:00")

    # Periodo 2: maggio
    ws.cell(row=14, column=1, value=date(2027, 5, 1))
    ws.cell(row=14, column=2, value=date(2027, 5, 31))
    ws.cell(row=14, column=3, value="6:15,7:00,8:30,10:00,11:30,17:30,18:30")
    ws.cell(row=14, column=4, value="6:15,7:00,8:30,10:00,11:30,17:30,18:30")

    return ws


# ======================================================================
# TEST — leggi_yaml
# ======================================================================


def test_leggi_yaml_valido(yaml_valido: Path) -> None:
    """Un YAML valido viene letto correttamente."""
    dati = leggi_yaml(yaml_valido)
    assert dati["nome"] == "Test"
    assert dati["valori"] == [1, 2, 3]


def test_leggi_yaml_vuoto(yaml_vuoto: Path) -> None:
    """Un YAML vuoto restituisce un dict vuoto."""
    dati = leggi_yaml(yaml_vuoto)
    assert dati == {}


def test_leggi_yaml_non_esiste(tmp_path: Path) -> None:
    """Un file inesistente solleva FileNotFoundError."""
    percorso = tmp_path / "non_esiste.yaml"
    with pytest.raises(FileNotFoundError):
        leggi_yaml(percorso)


def test_leggi_yaml_malformato(yaml_malformato: Path) -> None:
    """Uno YAML malformato solleva ValueError."""
    with pytest.raises(ValueError, match="YAML malformato"):
        leggi_yaml(yaml_malformato)


# ======================================================================
# TEST — parse_orari
# ======================================================================


def test_parse_orari_virgole() -> None:
    """Orari separati da virgola."""
    assert parse_orari("7:00,10:00,18:00") == ["7:00", "10:00", "18:00"]


def test_parse_orari_punto_virgola() -> None:
    """Orari separati da punto e virgola."""
    assert parse_orari("7:00;10:00") == ["7:00", "10:00"]


def test_parse_orari_spazi() -> None:
    """Orari separati da spazi."""
    assert parse_orari("7:00 10:00 18:00") == ["7:00", "10:00", "18:00"]


def test_parse_orari_vuoto() -> None:
    """Stringa vuota → lista vuota."""
    assert parse_orari("") == []


def test_parse_orari_valori_invalidi() -> None:
    """Valori non validi vengono ignorati."""
    assert parse_orari("7:00,pippo,18:00") == ["7:00", "18:00"]
    assert parse_orari("7.00,10:00") == ["10:00"]


def test_parse_orari_misto() -> None:
    """Separatori misti."""
    assert parse_orari("7:00, 10:00; 18:00") == ["7:00", "10:00", "18:00"]


# ======================================================================
# TEST — leggi_foglio_impostazioni
# ======================================================================


def test_leggi_impostazioni_valido(foglio_impostazioni_valido: Worksheet) -> None:
    """Un foglio valido viene letto correttamente."""
    dati = leggi_foglio_impostazioni(foglio_impostazioni_valido)
    assert dati["anno"] == 2027
    assert dati["nome_parrocchia"] == "San Pietro in Silki"
    assert dati["citta"] == "Sassari"
    assert len(dati["periodi"]) == 2
    assert dati["periodi"][0]["dal"] == date(2027, 10, 1)
    assert dati["periodi"][0]["orari_feriali"] == ["7:00", "10:00", "18:00"]
    assert dati["periodi"][1]["orari_festivi"] == [
        "6:15",
        "7:00",
        "8:30",
        "10:00",
        "11:30",
        "17:30",
        "18:30",
    ]


def test_leggi_impostazioni_anno_mancante() -> None:
    """Se manca l'anno → ValueError."""
    ws = crea_foglio_vuoto()
    ws["B5"] = None
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    with pytest.raises(ValueError, match="Anno"):
        leggi_foglio_impostazioni(ws)


def test_leggi_impostazioni_nome_mancante() -> None:
    """Se manca il nome parrocchia → ValueError."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = None
    ws["B7"] = "Test"

    with pytest.raises(ValueError, match="Nome parrocchia"):
        leggi_foglio_impostazioni(ws)


def test_leggi_impostazioni_riga_parziale() -> None:
    """Una riga con solo 'Dal' viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    ws.cell(row=13, column=1, value=date(2027, 10, 1))
    ws.cell(row=13, column=3, value="7:00")

    dati = leggi_foglio_impostazioni(ws)
    assert len(dati["periodi"]) == 0


def test_leggi_impostazioni_riga_non_data() -> None:
    """Una riga con valore non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws["B5"] = 2027
    ws["B6"] = "Test"
    ws["B7"] = "Test"

    ws.cell(row=13, column=1, value="01/10/2027")
    ws.cell(row=13, column=2, value="30/04/2028")

    dati = leggi_foglio_impostazioni(ws)
    assert len(dati["periodi"]) == 0

    # ======================================================================
    # FIXTURE — Foglio Intenzioni
    # ======================================================================


@pytest.fixture
def foglio_intenzioni_valido() -> Worksheet:
    """Crea un foglio Intenzioni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    # Riga 2: intenzione completa
    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value="Per la pace nel mondo")
    ws.cell(row=2, column=4, value=10.0)
    ws.cell(row=2, column=5, value=date(2027, 1, 20))
    ws.cell(row=2, column=6, value="Richiesta da Mario")

    # Riga 3: senza data applicazione (non ancora applicata)
    ws.cell(row=3, column=2, value=date(2027, 2, 1))
    ws.cell(row=3, column=3, value="Per i defunti della famiglia Rossi")
    ws.cell(row=3, column=4, value=20.0)

    # Riga 4: completamente vuota → ignorata

    # Riga 5: senza offerta
    ws.cell(row=5, column=2, value=date(2027, 3, 10))
    ws.cell(row=5, column=3, value="Per la salute di Anna")

    return ws


# ======================================================================
# TEST — leggi_foglio_intenzioni
# ======================================================================


def test_leggi_intenzioni_valido(foglio_intenzioni_valido: Worksheet) -> None:
    """Un foglio valido restituisce 3 intenzioni."""
    intenzioni = leggi_foglio_intenzioni(foglio_intenzioni_valido)
    assert len(intenzioni) == 3

    # Prima intenzione
    i1 = intenzioni[0]
    assert i1.numero == 1
    assert i1.data_consegna == date(2027, 1, 15)
    assert i1.testo == "Per la pace nel mondo"
    assert i1.offerta == 10.0
    assert i1.data_applicazione == date(2027, 1, 20)
    assert i1.note == "Richiesta da Mario"

    # Seconda: senza data applicazione
    i2 = intenzioni[1]
    assert i2.numero == 2
    assert i2.data_consegna == date(2027, 2, 1)
    assert i2.data_applicazione is None
    assert i2.offerta == 20.0

    # Terza: senza offerta
    i3 = intenzioni[2]
    assert i3.numero == 3
    assert i3.offerta is None


def test_leggi_intenzioni_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    intenzioni = leggi_foglio_intenzioni(ws)
    assert intenzioni == []


def test_leggi_intenzioni_riga_senza_testo() -> None:
    """Una riga con Data ma senza Intenzione viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    # Data compilata ma testo vuoto
    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value=None)

    # Riga valida
    ws.cell(row=3, column=2, value=date(2027, 2, 1))
    ws.cell(row=3, column=3, value="Intenzione valida")

    intenzioni = leggi_foglio_intenzioni(ws)
    assert len(intenzioni) == 1
    assert intenzioni[0].testo == "Intenzione valida"


def test_leggi_intenzioni_data_non_valida() -> None:
    """Una riga con Data non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    # Stringa invece di data
    ws.cell(row=2, column=2, value="15/01/2027")
    ws.cell(row=2, column=3, value="Testo")

    intenzioni = leggi_foglio_intenzioni(ws)
    assert intenzioni == []


def test_leggi_intenzioni_offerta_non_numerica() -> None:
    """Un'offerta non numerica viene impostata a None con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Intenzioni"

    ws.cell(row=2, column=2, value=date(2027, 1, 15))
    ws.cell(row=2, column=3, value="Testo valido")
    ws.cell(row=2, column=4, value="dieci euro")  # non numerico

    intenzioni = leggi_foglio_intenzioni(ws)
    assert len(intenzioni) == 1
    assert intenzioni[0].offerta is None


# ======================================================================
# FIXTURE — Foglio Matrimoni
# ======================================================================


@pytest.fixture
def foglio_matrimoni_valido() -> Worksheet:
    """Crea un foglio Matrimoni valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    # Riga 2: matrimonio completo
    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value="Maria Rossi e Luca Bianchi")
    ws.cell(row=2, column=4, value="333-1234567")
    ws.cell(row=2, column=5, value="Ricevimento in parrocchia")

    # Riga 3: senza contatti e note
    ws.cell(row=3, column=1, value=date(2028, 7, 15))
    ws.cell(row=3, column=2, value="10:00")
    ws.cell(row=3, column=3, value="Anna Verdi e Marco Neri")

    # Riga 4: completamente vuota → ignorata

    return ws


# ======================================================================
# TEST — leggi_foglio_matrimoni
# ======================================================================


def test_leggi_matrimoni_valido(foglio_matrimoni_valido: Worksheet) -> None:
    """Un foglio valido restituisce 2 matrimoni."""
    matrimoni = leggi_foglio_matrimoni(foglio_matrimoni_valido)
    assert len(matrimoni) == 2

    m1 = matrimoni[0]
    assert m1.data == date(2028, 6, 10)
    assert m1.ora == "15:30"
    assert m1.nome_sposi == "Maria Rossi e Luca Bianchi"
    assert m1.contatti == "333-1234567"
    assert m1.note == "Ricevimento in parrocchia"

    m2 = matrimoni[1]
    assert m2.data == date(2028, 7, 15)
    assert m2.ora == "10:00"
    assert m2.contatti == ""
    assert m2.note == ""


def test_leggi_matrimoni_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_senza_ora() -> None:
    """Una riga senza Ora viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value=None)
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_ora_non_valida() -> None:
    """Una riga con Ora malformata viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15.30")  # punto invece di due punti
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_senza_nome() -> None:
    """Una riga senza Nome sposi viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value=None)

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_data_non_data() -> None:
    """Una riga con Data non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value="10/06/2028")
    ws.cell(row=2, column=2, value="15:30")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_matrimoni(ws) == []


def test_leggi_matrimoni_ora_nativa() -> None:
    """Un'ora nativa (datetime.time) viene convertita in stringa HH:MM."""
    from datetime import time

    ws = crea_foglio_vuoto()
    ws.title = "Matrimoni"

    ws.cell(row=2, column=1, value=date(2028, 6, 10))
    ws.cell(row=2, column=2, value=time(15, 30))
    ws.cell(row=2, column=3, value="Test")

    matrimoni = leggi_foglio_matrimoni(ws)
    assert len(matrimoni) == 1
    assert matrimoni[0].ora == "15:30"


# ======================================================================
# FIXTURE — Foglio Note
# ======================================================================


@pytest.fixture
def foglio_note_valido() -> Worksheet:
    """Crea un foglio Note valido in memoria."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    # Riga 2: periodo multi-giorno
    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value=date(2027, 3, 16))
    ws.cell(row=2, column=3, value="Triduo San Salvatore")

    # Riga 3: giorno singolo (Al vuoto → Al = Dal)
    ws.cell(row=3, column=1, value=date(2027, 5, 30))
    ws.cell(row=3, column=3, value="Corpus Domini, vietato pomeriggio")

    # Riga 4: completamente vuota → ignorata

    # Riga 5: periodo lungo
    ws.cell(row=5, column=1, value=date(2027, 5, 1))
    ws.cell(row=5, column=2, value=date(2027, 5, 31))
    ws.cell(row=5, column=3, value="Mese mariano, programma da definire")

    return ws


# ======================================================================
# TEST — leggi_foglio_note
# ======================================================================


def test_leggi_note_valido(foglio_note_valido: Worksheet) -> None:
    """Un foglio valido restituisce 3 note."""
    note = leggi_foglio_note(foglio_note_valido)
    assert len(note) == 3

    # Prima: multi-giorno
    n1 = note[0]
    assert n1["dal"] == date(2027, 3, 14)
    assert n1["al"] == date(2027, 3, 16)
    assert n1["nota"] == "Triduo San Salvatore"

    # Seconda: giorno singolo (Al = Dal)
    n2 = note[1]
    assert n2["dal"] == date(2027, 5, 30)
    assert n2["al"] == date(2027, 5, 30)
    assert n2["nota"] == "Corpus Domini, vietato pomeriggio"

    # Terza: mese intero
    n3 = note[2]
    assert n3["dal"] == date(2027, 5, 1)
    assert n3["al"] == date(2027, 5, 31)


def test_leggi_note_foglio_vuoto() -> None:
    """Un foglio vuoto restituisce lista vuota."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    assert leggi_foglio_note(ws) == []


def test_leggi_note_senza_nota() -> None:
    """Una riga con Dal ma senza Nota viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=3, value=None)

    assert leggi_foglio_note(ws) == []


def test_leggi_note_dal_non_data() -> None:
    """Una riga con Dal non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value="14/03/2027")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_non_data() -> None:
    """Una riga con Al non-data viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value="16/03/2027")
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_prima_di_dal() -> None:
    """Una riga con Al < Dal viene ignorata con warning."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 16))
    ws.cell(row=2, column=2, value=date(2027, 3, 14))  # Al < Dal
    ws.cell(row=2, column=3, value="Test")

    assert leggi_foglio_note(ws) == []


def test_leggi_note_al_uguale_dal() -> None:
    """Una nota con Al = Dal è valida."""
    ws = crea_foglio_vuoto()
    ws.title = "Note"

    ws.cell(row=2, column=1, value=date(2027, 3, 14))
    ws.cell(row=2, column=2, value=date(2027, 3, 14))
    ws.cell(row=2, column=3, value="Test")

    note = leggi_foglio_note(ws)
    assert len(note) == 1
    assert note[0]["dal"] == note[0]["al"]


# ======================================================================
# TEST — unisci_config
# ======================================================================


def test_unisci_config_minimo() -> None:
    """Unione con dati minimi."""
    dati_yaml = {
        "nome_parrocchia": "Test",
        "citta": "Roma",
        "festivi_fissi": [],
        "eccezioni_festivi": [],
        "divieti_pomeridiani": {"fissi": [], "mobili": []},
        "ricorrenze_proprie": [],
    }
    dati_excel = {
        "anno": 2027,
        "nome_parrocchia": "Test",
        "citta": "Roma",
        "periodi": [],
    }

    config = unisci_config(dati_yaml, dati_excel)
    assert config.nome_parrocchia == "Test"
    assert config.citta == "Roma"
    assert config.anno == 2027
    assert config.periodi == []


def test_unisci_config_feste() -> None:
    """Le feste del YAML diventano oggetti Festa."""
    dati_yaml = {
        "nome_parrocchia": "Test",
        "citta": "Roma",
        "festivi_fissi": [
            {"mese": 12, "giorno": 25, "nome": "Natale"},
        ],
        "eccezioni_festivi": [],
        "divieti_pomeridiani": {"fissi": [], "mobili": []},
        "ricorrenze_proprie": [],
    }
    dati_excel = {"anno": 2027, "periodi": []}

    config = unisci_config(dati_yaml, dati_excel)
    assert len(config.festivi_fissi) == 1
    assert config.festivi_fissi[0].nome == "Natale"
    assert config.festivi_fissi[0].mese == 12
    assert config.festivi_fissi[0].giorno == 25


def test_unisci_config_ricorrenza_fissa() -> None:
    """Una ricorrenza a data fissa diventa oggetto Ricorrenza."""
    dati_yaml = {
        "nome_parrocchia": "Test",
        "citta": "Roma",
        "festivi_fissi": [],
        "eccezioni_festivi": [],
        "divieti_pomeridiani": {"fissi": [], "mobili": []},
        "ricorrenze_proprie": [
            {
                "nome": "San Salvatore",
                "data": {"mese": 3, "giorno": 17},
                "durata_giorni": 2,
            },
        ],
    }
    dati_excel = {"anno": 2027, "periodi": []}

    config = unisci_config(dati_yaml, dati_excel)
    assert len(config.ricorrenze_proprie) == 1
    ric = config.ricorrenze_proprie[0]
    assert ric.nome == "San Salvatore"
    assert ric.data_inizio == date(2027, 3, 17)
    assert ric.data_fine == date(2027, 3, 18)


def test_unisci_config_mismatch_nome() -> None:
    """Nome YAML vince su nome Excel."""
    dati_yaml = {"nome_parrocchia": "YAML Parrocchia", "citta": "Roma"}
    dati_excel = {"anno": 2027, "nome_parrocchia": "Excel Parrocchia", "citta": "Roma"}

    config = unisci_config(dati_yaml, dati_excel)
    assert config.nome_parrocchia == "YAML Parrocchia"


def test_unisci_config_periodi() -> None:
    """I periodi Excel diventano oggetti Periodo."""
    dati_yaml = {"nome_parrocchia": "Test", "citta": "Roma"}
    dati_excel = {
        "anno": 2027,
        "periodi": [
            {
                "dal": date(2027, 1, 1),
                "al": date(2027, 1, 31),
                "orari_feriali": ["7:00"],
                "orari_festivi": ["10:00"],
            },
        ],
    }

    config = unisci_config(dati_yaml, dati_excel)
    assert len(config.periodi) == 1
    assert config.periodi[0].dal == date(2027, 1, 1)
    assert config.periodi[0].orari_feriali == ["7:00"]


# ======================================================================
# TEST — carica_config
# ======================================================================


def test_carica_config_profilo_inesistente(tmp_path: Path) -> None:
    """Un profilo inesistente solleva FileNotFoundError."""
    with pytest.raises(FileNotFoundError, match="Profilo non trovato"):
        carica_config("profilo_inesistente", cartella_configs=tmp_path)


def test_carica_config_yaml_mancante(tmp_path: Path) -> None:
    """Un profilo senza regole.yaml solleva FileNotFoundError."""
    profilo_dir = tmp_path / "test_profilo"
    profilo_dir.mkdir()

    with pytest.raises(FileNotFoundError, match="regole.yaml non trovato"):
        carica_config("test_profilo", cartella_configs=tmp_path)


def test_carica_config_excel_mancante(tmp_path: Path) -> None:
    """Un profilo senza config.xlsx solleva FileNotFoundError."""
    profilo_dir = tmp_path / "test_profilo"
    profilo_dir.mkdir()

    # Crea solo regole.yaml
    (profilo_dir / "regole.yaml").write_text(
        'nome_parrocchia: "Test"\ncitta: "Roma"\n',
        encoding="utf-8",
    )

    with pytest.raises(FileNotFoundError, match="config.xlsx non trovato"):
        carica_config("test_profilo", cartella_configs=tmp_path)
